"""Write learning progress into Marathon_Daily_Planner_v4.xlsx.

Run with:  uv run --no-project --with openpyxl python tools/update_planner.py <command> ...

Commands:
  larp L01 --scores 2,1,2,0,1,2,2,1,1   score questions 1..N (0-2); score < 2 ticks 'Revisit'
  lab LAB00 [--note "one-line explain-it note"]
  pattern "Rate limiting" [--theory] [--built] [--explain] [--design]
  note 2026-10-28 "LARP L01 done on 29 Oct (catch-up)"   append to Daily Plan 'Notes / carry-over'
  done S01 "IP, ports, packets" [--score 5/6] [--note "struggled with 0.0.0.0"]
                                                         log a finished session in the 'Course Log' sheet
  next "S02b:Bytes vs text" "S03:Sockets hands-on"       replace the 'up next' rows in 'Course Log'
  show                                                   print what's ticked so far
"""
import argparse
import datetime as dt
import sys
from pathlib import Path

import openpyxl

PLANNER = Path(__file__).resolve().parent.parent / "Marathon_Daily_Planner_v4.xlsx"
TICK, EMPTY = "☑", "☐"


def header_map(ws, row: int) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[row] if c.value is not None}


def cmd_larp(wb, args) -> None:
    ws = wb["LARP Bank"]
    cols = header_map(ws, 4)
    scores = [int(s) for s in args.scores.split(",")] if args.scores else []
    if any(s not in (0, 1, 2) for s in scores):
        sys.exit("scores must be 0, 1 or 2")
    found = 0
    for row in ws.iter_rows(min_row=5):
        if row[cols["ID"] - 1].value != args.id:
            continue
        num = row[cols["#"] - 1].value
        if not isinstance(num, int) or num > len(scores):
            continue
        score = scores[num - 1]
        ws.cell(row[0].row, cols["My score 0-2"]).value = score
        ws.cell(row[0].row, cols["Revisit"]).value = TICK if score < 2 else EMPTY
        found += 1
    if not found:
        sys.exit(f"no scorable questions found for {args.id}")
    print(f"{args.id}: scored {found} questions, total {sum(scores)}/{2 * len(scores)}")


def cmd_lab(wb, args) -> None:
    ws = wb["Backend Labs"]
    cols = header_map(ws, 4)
    for row in ws.iter_rows(min_row=5):
        lab = str(row[cols["Lab"] - 1].value or "")
        if lab.split(" ")[0] == args.id:
            ws.cell(row[0].row, cols["Done"]).value = TICK
            if args.note:
                ws.cell(row[0].row, cols["Explain-it note (1 line)"]).value = args.note
            print(f"ticked {lab}")
            return
    sys.exit(f"lab {args.id} not found")


def cmd_pattern(wb, args) -> None:
    ws = wb["Backend Patterns"]
    cols = header_map(ws, 4)
    flags = {"Theory": args.theory, "Built": args.built, "Explain": args.explain, "Design": args.design}
    matches = [r for r in ws.iter_rows(min_row=5) if args.name.lower() in str(r[cols["Pattern"] - 1].value or "").lower()]
    if len(matches) != 1:
        names = [r[cols["Pattern"] - 1].value for r in matches]
        sys.exit(f"'{args.name}' matched {len(matches)} patterns: {names}")
    row = matches[0]
    for col_name, on in flags.items():
        if on:
            ws.cell(row[0].row, cols[col_name]).value = TICK
    print(f"updated {row[cols['Pattern'] - 1].value}: {[k for k, v in flags.items() if v]}")


def cmd_note(wb, args) -> None:
    ws = wb["Daily Plan"]
    cols = header_map(ws, 1)
    day = dt.date.fromisoformat(args.date)
    for row in ws.iter_rows(min_row=2):
        value = row[cols["Date"] - 1].value
        if isinstance(value, dt.datetime) and value.date() == day:
            cell = ws.cell(row[0].row, cols["Notes / carry-over"])
            cell.value = f"{cell.value}\n{args.text}" if cell.value else args.text
            print(f"note added to {args.date}")
            return
    sys.exit(f"date {args.date} not in Daily Plan")


LOG_SHEET = "Course Log"
LOG_HEADERS = ["ID", "Topic", "Status", "Date", "Score", "Struggles / insight"]


def course_log(wb):
    if LOG_SHEET in wb.sheetnames:
        return wb[LOG_SHEET]
    ws = wb.create_sheet(LOG_SHEET, index=2)  # right after 'Today'
    ws.append(LOG_HEADERS)
    for col, width in zip("ABCDEF", (8, 40, 8, 12, 8, 60)):
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    return ws


def log_rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(min_row=2, max_col=len(LOG_HEADERS), values_only=True) if r[0]]


def write_log(ws, rows: list[list]) -> None:
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        ws.append(r)


def cmd_done(wb, args) -> None:
    ws = course_log(wb)
    rows = [r for r in log_rows(ws) if not (r[0] == args.id and r[2] == "next")]
    done = [r for r in rows if r[2] == "done"]
    upcoming = [r for r in rows if r[2] != "done"]
    done.append([args.id, args.topic, "done", dt.date.today().isoformat(), args.score or "", args.note or ""])
    write_log(ws, done + upcoming)
    print(f"logged {args.id} as done")


def cmd_next(wb, args) -> None:
    ws = course_log(wb)
    done = [r for r in log_rows(ws) if r[2] == "done"]
    upcoming = []
    for item in args.items:
        sid, _, topic = item.partition(":")
        upcoming.append([sid.strip(), topic.strip(), "next", "", "", ""])
    write_log(ws, done + upcoming)
    print(f"up next: {[u[0] for u in upcoming]}")


def cmd_show(wb, _args) -> None:
    ws = wb["LARP Bank"]
    cols = header_map(ws, 4)
    scored: dict[str, list[int]] = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        score = row[cols["My score 0-2"] - 1]
        if isinstance(score, int):
            scored.setdefault(row[cols["ID"] - 1], []).append(score)
    for lid, s in scored.items():
        print(f"LARP {lid}: {sum(s)}/{2 * len(s)}")
    for row in wb["Backend Labs"].iter_rows(min_row=5, values_only=True):
        if row[4] == TICK:
            print(f"LAB done: {row[2]}")
    pcols = header_map(wb["Backend Patterns"], 4)
    for row in wb["Backend Patterns"].iter_rows(min_row=5, values_only=True):
        ticks = [k for k in ("Theory", "Built", "Explain", "Design") if row[pcols[k] - 1] == TICK]
        if ticks:
            print(f"Pattern {row[pcols['Pattern'] - 1]}: {ticks}")


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("larp"); s.add_argument("id"); s.add_argument("--scores", required=True)
    s = sub.add_parser("lab"); s.add_argument("id"); s.add_argument("--note")
    s = sub.add_parser("pattern"); s.add_argument("name")
    for f in ("theory", "built", "explain", "design"):
        s.add_argument(f"--{f}", action="store_true")
    s = sub.add_parser("note"); s.add_argument("date"); s.add_argument("text")
    s = sub.add_parser("done"); s.add_argument("id"); s.add_argument("topic")
    s.add_argument("--score"); s.add_argument("--note")
    s = sub.add_parser("next"); s.add_argument("items", nargs="+", help='"ID:topic" pairs')
    sub.add_parser("show")
    args = p.parse_args()

    wb = openpyxl.load_workbook(PLANNER)
    commands = {"larp": cmd_larp, "lab": cmd_lab, "pattern": cmd_pattern, "note": cmd_note,
                "done": cmd_done, "next": cmd_next, "show": cmd_show}
    commands[args.cmd](wb, args)
    if args.cmd != "show":
        try:
            wb.save(PLANNER)
        except PermissionError:
            sys.exit(f"Can't save {PLANNER.name}: close it in Excel and run again.")


if __name__ == "__main__":
    main()
